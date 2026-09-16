from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
PPTX = OUT / "Agentic_Navigation_Report.pptx"
COVER_IMG = OUT / "cover.jpg"
PPT_DEMO_VIDEO = OUT / "ppt_demo.mp4"
PPT_DEMO_COVER = OUT / "cover.jpg"
CLOSED_LOOP_VIDEO = OUT / "output_closed_loop" / "episode=86-ckpt=0-spl=0.00.mp4"
CLOSED_LOOP_FRAME = OUT / "output_closed_loop" / "demo_frame.jpg"

DARK = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
LIGHT = RGBColor(0xEC, 0xF3, 0xFA)
GRAY = RGBColor(0x5A, 0x64, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xE8, 0x8B, 0x2E)


def set_run(run, text, size=18, bold=False, color=DARK, italic=False):
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("a:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("a:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("a:ea"), "Microsoft YaHei")
    rFonts.set(qn("a:cs"), "Microsoft YaHei")


def add_textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        set_run(run, text, size=size, bold=bold, color=color)
    return box


def add_title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.12)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, title, size=26, bold=True, color=WHITE)
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        set_run(r2, subtitle, size=14, bold=False, color=RGBColor(0xCF, 0xE0, 0xF2))


def add_bullets(slide, items, left=0.7, top=1.4, width=12.0, height=5.4, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        set_run(run, "▪ " + item, size=size, color=DARK)
    return box


def add_section_slide(prs, number, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK
    bg.line.fill.background()
    add_textbox(slide, 1.0, 2.5, 11.3, 2.5, [
        (f"{number}", 48, True, ORANGE),
        (title, 40, True, WHITE),
        (subtitle, 20, False, RGBColor(0xCF, 0xE0, 0xF2)),
    ], align=PP_ALIGN.CENTER)
    return slide


def add_shape_box(slide, x, y, w, h, text, fill=BLUE, font_color=WHITE, size=18, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, text, size=size, bold=bold, color=font_color)
    return shape


def add_arrow(slide, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRAY
    shape.line.fill.background()
    return shape


def add_table(slide, rows, cols, left, top, width, height, data, col_widths=None):
    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c > 0 else PP_ALIGN.CENTER
                for run in p.runs:
                    set_run(run, run.text, size=12 if r else 13, bold=(r == 0), color=WHITE if r == 0 else DARK)
    return table


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. 封面
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK
    bg.line.fill.background()
    add_textbox(slide, 0.9, 1.4, 8.2, 4.2, [
        ("Agentic Navigation", 54, True, WHITE),
        ("轻量开源 Agent 框架与 NaVILA 协作", 34, True, RGBColor(0x9C, 0xC7, 0xFF)),
        ("组会汇报 · 2026-09-04", 20, False, RGBColor(0xCF, 0xE0, 0xF2)),
    ], align=PP_ALIGN.LEFT)
    if COVER_IMG.exists():
        slide.shapes.add_picture(str(COVER_IMG), Inches(8.2), Inches(1.5), Inches(4.4), Inches(2.5))
    add_textbox(slide, 0.9, 6.4, 11.5, 0.7, [
        ("Agno Supervisor  ×  NaVILA Executor  ×  Habitat/Robot", 18, True, ORANGE),
    ], align=PP_ALIGN.LEFT)

    # 2. 项目背景与目标
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "项目背景与目标")
    add_bullets(slide, [
        "长期目标：将 VLN 部署到机器狗，实现跨楼宇的自主导航。",
        "已有基础：NaVILA 视觉-语言-动作模型已本地部署，Habitat/VLN-CE 仿真环境可用。",
        "核心缺口：缺少一个负责长程规划、记忆、工具调用和异常重规划的高层 Agent 编排层。",
        "本次工作：调研轻量开源 Agent 框架，确定 Agno，并实现 Agno × NaVILA 的真实闭环演示。",
    ])

    # 3. Agent 框架在项目中的作用
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "Agent 框架在项目中的作用")
    add_textbox(slide, 0.7, 1.25, 12.0, 0.6, [
        ("Agent 是“任务大脑”，负责慢思考与全局协调；NaVILA 负责快行动。", 20, True, BLUE),
    ])
    roles = [
        ("任务理解与规划", "将“去 5 楼会议室”分解为可执行子目标"),
        ("外部信息整合", "地图 API、POI、楼层语义图、电梯/门禁状态"),
        ("Skill / Tool 编排", "复用导航 SOP，调用地图、机器人、人工求助等工具"),
        ("记忆管理", "会话记忆、用户偏好、成功路线与失败经验"),
        ("异常监督与重规划", "监听 blocked / failed / safety_stop 并重规划"),
        ("安全与人机协同", "高风险动作确认、人工接管"),
    ]
    for i, (a, b) in enumerate(roles):
        x = 0.7 + (i % 2) * 6.2
        y = 2.0 + (i // 2) * 1.6
        add_shape_box(slide, x, y, 5.9, 1.2, "", fill=LIGHT, font_color=DARK, size=14, bold=False)
        add_textbox(slide, x + 0.1, y + 0.08, 5.7, 1.0, [
            (a, 16, True, BLUE),
            (b, 14, False, DARK),
        ])

    # 4. 为什么需要 Agent，而不是只用 NaVILA
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "为什么需要 Agent，而不是只用 NaVILA")
    add_bullets(slide, [
        "NaVILA 是两层 VLA：VLM 生成中层动作 + locomotion policy 执行；适合短程、视觉驱动导航。",
        "跨楼宇任务需要长期记忆、路线规划、外部工具调用和失败恢复，这超出 NaVILA 原生能力。",
        "NaVILA 的输入应保持简短：当前子目标 + 历史帧，不应塞入整张地图和长期对话。",
        "Agent 不应进入高频控制循环；它只在新任务、子目标完成、异常事件时介入。",
    ])

    # 5. 调研需求与评估标准
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "调研需求与评估标准")
    criteria = [
        ("轻量", "启动快，可嵌入导航运行时"),
        ("开源许可", "可自托管、可商用改造"),
        ("多模态", "原生 image/video/audio 输入"),
        ("记忆", "短期会话 + 长期用户/任务记忆，可持久化"),
        ("Skill", "支持 SKILL.md 一类可复用流程"),
        ("Tool/MCP", "自定义工具 + 标准 MCP 接入"),
        ("部署", "易以 FastAPI/Docker 服务化"),
        ("模型无关", "可接 DeepSeek、Qwen、Ollama、vLLM 等"),
    ]
    for i, (a, b) in enumerate(criteria):
        x = 0.7 + (i % 2) * 6.2
        y = 1.5 + (i // 2) * 1.25
        add_shape_box(slide, x, y, 5.9, 1.0, "", fill=LIGHT, font_color=DARK, size=14, bold=False)
        add_textbox(slide, x + 0.15, y + 0.1, 5.6, 0.9, [
            (a, 16, True, BLUE),
            (b, 14, False, DARK),
        ])

    # 6. 候选框架总览
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "候选框架总览")
    add_textbox(slide, 0.7, 1.4, 12.0, 4.8, [
        ("通用轻量多模态：", 20, True, BLUE),
        ("Agno、OpenAI Agents SDK、LightAgent", 20, False, DARK),
        ("记忆优先：", 20, True, BLUE),
        ("Letta（MemGPT）", 20, False, DARK),
        ("编排/多智能体：", 20, True, BLUE),
        ("LangGraph、CrewAI、AutoGen/AG2", 20, False, DARK),
        ("具身/GUI 参考：", 20, True, BLUE),
        ("HoloAgent、Cortex、Agent TARS、browser-use", 20, False, DARK),
    ], align=PP_ALIGN.LEFT)

    # 7. 横向对比
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "候选框架横向对比")
    data = [
        ["框架", "轻量", "多模态", "记忆", "Skill", "Tool/MCP", "部署", "适配度"],
        ["Agno", "极轻", "原生", "强", "原生", "函数+MCP", "AgentOS/Docker", "★★★★★"],
        ["OpenAI Agents SDK", "轻", "API 支持", "Sessions", "有", "函数+MCP", "自建", "★★★★☆"],
        ["Letta", "中", "部分", "极强", "无原生", "工具+部分MCP", "Server+Docker", "★★★★☆"],
        ["LangGraph", "重", "依赖集成", "Checkpoint", "无原生", "生态大+MCP", "Platform", "★★★☆☆"],
        ["CrewAI", "中重", "有限", "统一记忆", "有", "工具+MCP", "自建", "★★★☆☆"],
        ["LightAgent", "轻", "较弱", "可插拔", "原生", "工具+MCP", "自建", "★★★★☆"],
    ]
    add_table(slide, 7, 8, 0.4, 1.4, 12.5, 4.6, data, col_widths=[2.4, 1.0, 1.4, 1.6, 1.2, 1.9, 1.4, 1.4])

    # 8. 为什么选择 Agno
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "为什么选择 Agno")
    add_bullets(slide, [
        "轻量：单 Agent 实例化开销小，适合嵌入导航任务。",
        "原生多模态：Image / Audio / Video / File 统一输入。",
        "成熟记忆：会话 + 用户记忆，SQLite/Postgres 持久化，支持 automatic / agentic。",
        "原生 SKILL.md：Anthropic 风格技能包，可渐进加载 scripts/references。",
        "工具生态：自定义函数工具 + 100+ Toolkit + MCPTools。",
        "部署友好：AgentOS（FastAPI）+ Docker，便于服务化。",
        "风险：MPL-2.0 许可、API 迭代快需锁版本；视频输入模型支持有限，可走图像序列。",
    ])

    # 9. Agno × NaVILA 协作架构
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "Agno × NaVILA：Supervisor–Executor")
    add_shape_box(slide, 4.0, 1.2, 5.3, 0.9, "Agno Supervisor\n规划 / 重规划 / 记忆 / 工具", fill=DARK)
    add_arrow(slide, 6.25, 2.1, 0.8, 0.7)
    add_shape_box(slide, 3.2, 2.8, 6.9, 1.5, "VLAExecutor（NaVILA）\nframe ring buffer → VLA → mid-level action", fill=BLUE)
    add_arrow(slide, 6.25, 4.3, 0.8, 0.7)
    add_shape_box(slide, 4.0, 5.0, 5.3, 0.9, "Robot / Habitat\n实时执行 + 安全监控", fill=GRAY)
    add_textbox(slide, 0.7, 6.2, 12.0, 0.8, [
        ("Event: completed / failed / blocked / uncertain / safety_stop / need_help", 14, True, ORANGE),
    ], align=PP_ALIGN.CENTER)

    # 10. 控制回路与职责边界
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "控制回路与职责边界")
    data = [
        ["回路", "角色", "频率", "触发"],
        ["任务/重规划", "Agno", "0.1–1 Hz", "新任务、子目标完成、异常"],
        ["中层动作", "NaVILA", "1–5 Hz", "当前子目标未完成"],
        ["底层运动", "Robot/Controller", "50–200 Hz", "独立实时闭环"],
        ["安全", "Safety Monitor", "独立高频", "急停、碰撞、失稳"],
    ]
    add_table(slide, 5, 4, 0.7, 1.5, 12.0, 3.2, data, col_widths=[3.0, 3.5, 2.5, 3.0])
    add_textbox(slide, 0.7, 5.3, 12.0, 1.0, [
        ("核心原则：Agent 不进高频控制循环，NaVILA 不直接理解长程任务。", 18, True, BLUE),
    ], align=PP_ALIGN.CENTER)

    # 11. 演示：Agno 与 NaVILA 协作导航
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "演示：Agno 高层规划 → NaVILA 执行 → 事件回传")
    demo_video = PPT_DEMO_VIDEO if PPT_DEMO_VIDEO.exists() else CLOSED_LOOP_VIDEO
    demo_cover = PPT_DEMO_COVER if PPT_DEMO_COVER.exists() else CLOSED_LOOP_FRAME
    if demo_video.exists():
        slide.shapes.add_movie(
            str(demo_video),
            Inches(0.7), Inches(1.4), Inches(8.0), Inches(4.5),
            poster_frame_image=str(demo_cover) if demo_cover.exists() else None,
            mime_type="video/mp4",
        )
    elif demo_cover.exists():
        slide.shapes.add_picture(str(demo_cover), Inches(0.7), Inches(1.4), Inches(8.0), Inches(4.5))
    add_textbox(slide, 9.0, 1.6, 3.8, 4.6, [
        ("运行链路", 20, True, BLUE),
        ("Agno /plan", 16, False, DARK),
        ("→ SubGoals", 14, False, GRAY),
        ("NaVILA /navigate", 16, False, DARK),
        ("→ mid-level action", 14, False, GRAY),
        ("Habitat step", 16, False, DARK),
        ("→ 新观测帧", 14, False, GRAY),
        ("event → Agno replan", 16, False, DARK),
    ], align=PP_ALIGN.LEFT)

    # 12. 当前进展与下一步
    slide = prs.slides.add_slide(blank)
    add_title_bar(slide, "当前进展与下一步")
    add_textbox(slide, 0.7, 1.3, 5.8, 3.2, [
        ("已完成", 22, True, BLUE),
        ("· 框架调研与 Agno 选型", 18, False, DARK),
        ("· Supervisor–Executor 脚手架", 18, False, DARK),
        ("· PPT 演示视频生成", 18, False, DARK),
        ("· 真实 Habitat 闭环跑通", 18, False, DARK),
    ])
    add_textbox(slide, 7.0, 1.3, 5.8, 3.2, [
        ("下一步", 22, True, ORANGE),
        ("· 子目标完成判据，提升成功率", 18, False, DARK),
        ("· 地图 API + 室内语义图", 18, False, DARK),
        ("· ROS2 真机接入", 18, False, DARK),
        ("· Postgres/向量记忆库", 18, False, DARK),
    ])

    prs.save(PPTX)
    print(f"PPT 已生成：{PPTX}")


if __name__ == "__main__":
    build()
